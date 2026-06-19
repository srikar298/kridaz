"""
Kridaz — Marketing Feature Documentation Generator
Generates a professional .docx for the marketing & design team.
"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

# ──────────────────────────────────────────────────────────────
# Color palette
# ──────────────────────────────────────────────────────────────
KRIDAZ_GREEN = RGBColor(0x00, 0xC1, 0x87)
DARK_BG = RGBColor(0x05, 0x05, 0x05)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
MEDIUM_GRAY = RGBColor(0x66, 0x66, 0x66)
DARK_TEXT = RGBColor(0x1A, 0x1A, 0x1A)
SECTION_BLUE = RGBColor(0x0A, 0x84, 0xFF)
SECTION_ORANGE = RGBColor(0xFF, 0x8C, 0x00)
SECTION_PURPLE = RGBColor(0x8B, 0x5C, 0xF6)
SECTION_RED = RGBColor(0xEF, 0x44, 0x44)


def set_cell_shading(cell, color_hex):
    """Set background color for a table cell."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_border(cell, **kwargs):
    """Set cell borders. kwargs like top=("single","4","000000")"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge, val in kwargs.items():
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), val[0])
        element.set(qn("w:sz"), val[1])
        element.set(qn("w:color"), val[2])
        element.set(qn("w:space"), "0")
        tcBorders.append(element)
    tcPr.append(tcBorders)


def add_horizontal_rule(doc):
    """Add a thin horizontal line separator."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:color"), "CCCCCC")
    bottom.set(qn("w:space"), "1")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_feature_section(doc, number, feature):
    """Add a complete feature section to the document."""

    # ── Feature heading ──
    heading = doc.add_heading(level=2)
    run = heading.add_run(f"Feature #{number}: {feature['name']}")
    run.font.color.rgb = DARK_TEXT
    run.font.size = Pt(16)

    # ── Tagline ──
    tagline_p = doc.add_paragraph()
    tagline_p.paragraph_format.space_before = Pt(0)
    tagline_p.paragraph_format.space_after = Pt(6)
    run = tagline_p.add_run(f"✦ \"{feature['tagline']}\"")
    run.font.color.rgb = KRIDAZ_GREEN
    run.font.size = Pt(12)
    run.bold = True

    # ── Info table ──
    info_table = doc.add_table(rows=0, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.LEFT

    info_fields = [
        ("🎯 Motto", feature["motto"]),
        ("👤 Target Audience", feature["audience"]),
        ("📌 Category", feature["category"]),
    ]

    for label, value in info_fields:
        row = info_table.add_row()
        cell_label = row.cells[0]
        cell_value = row.cells[1]
        cell_label.width = Cm(4)
        cell_value.width = Cm(12)

        p = cell_label.paragraphs[0]
        run = p.add_run(label)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = DARK_TEXT

        p = cell_value.paragraphs[0]
        run = p.add_run(value)
        run.font.size = Pt(10)
        run.font.color.rgb = MEDIUM_GRAY

        set_cell_shading(cell_label, "F8F8F8")
        border_args = {
            "top": ("single", "2", "E0E0E0"),
            "bottom": ("single", "2", "E0E0E0"),
            "start": ("single", "2", "E0E0E0"),
            "end": ("single", "2", "E0E0E0"),
        }
        set_cell_border(cell_label, **border_args)
        set_cell_border(cell_value, **border_args)

    doc.add_paragraph("")

    # ── What It Does ──
    p = doc.add_paragraph()
    run = p.add_run("📖 What It Does")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_TEXT
    p = doc.add_paragraph(feature["what_it_does"])
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        run.font.size = Pt(10)

    # ── How the User Uses It ──
    p = doc.add_paragraph()
    run = p.add_run("🚶 How the User Uses It (Step-by-Step Flow)")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_TEXT

    for i, step in enumerate(feature["user_flow"], 1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(step)
        run.font.size = Pt(10)

    doc.add_paragraph("")

    # ── Purpose & Value ──
    p = doc.add_paragraph()
    run = p.add_run("💡 Purpose & Business Value")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_TEXT
    p = doc.add_paragraph(feature["purpose"])
    p.paragraph_format.space_after = Pt(8)
    for run in p.runs:
        run.font.size = Pt(10)

    # ── Key Highlights ──
    if feature.get("highlights"):
        p = doc.add_paragraph()
        run = p.add_run("⭐ Key Highlights for Marketing Posts")
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = DARK_TEXT

        for highlight in feature["highlights"]:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(highlight)
            run.font.size = Pt(10)

    doc.add_paragraph("")

    # ── Visual Design Direction ──
    p = doc.add_paragraph()
    run = p.add_run("🎨 Visual Design Direction for Graphics")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = KRIDAZ_GREEN

    design = feature["design_direction"]
    design_fields = [
        ("Suggested Style", design["style"]),
        ("Color Palette", design["colors"]),
        ("Imagery / Icons", design["imagery"]),
        ("Post Formats", design["formats"]),
        ("Suggested CTA", design["cta"]),
    ]
    for label, value in design_fields:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(f"{label}: ")
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = DARK_TEXT
        run = p.add_run(value)
        run.font.size = Pt(10)
        run.font.color.rgb = MEDIUM_GRAY

    add_horizontal_rule(doc)


# ──────────────────────────────────────────────────────────────
# FEATURE DATA — Deep-dived from the codebase
# ──────────────────────────────────────────────────────────────
FEATURES = [
    # ═══════════════════════════════════════════════════════════
    # PERSONA 1: PLAYER (USER)
    # ═══════════════════════════════════════════════════════════
    {
        "name": "Venue Discovery & Booking",
        "tagline": "Find. Book. Play — in under 60 seconds.",
        "motto": "Your next game is just a tap away.",
        "audience": "Players, casual sports enthusiasts, weekend warriors",
        "category": "🏏 Player Features — Core",
        "what_it_does": "Users can discover nearby sports venues (turfs/grounds) with real-time availability, filter by sport type, price, location, and time. They can view venue details including photos, amenities, reviews, and ratings. The booking flow lets users select date, time slots, and pay securely via Razorpay or Kridaz Wallet. After booking, they receive a digital Booking Pass and a downloadable Invoice.",
        "user_flow": [
            "Opens the app → taps 'Venues' in the bottom navigation bar.",
            "Browses venues on a map or list view. Filters by sport, price range, distance.",
            "Taps a venue card → sees full details: photos carousel, price/hr, open hours, amenities, reviews.",
            "Taps 'Book Now' → selects date and available time slots on a visual calendar grid.",
            "Proceeds to Checkout → chooses payment method (Razorpay / Wallet). Applies coupon if available.",
            "Confirms payment → receives a Booking Pass (shareable digital card) and a PDF Invoice.",
            "On game day, shows the Booking Pass at the venue for entry."
        ],
        "purpose": "This is the bread-and-butter revenue feature of Kridaz. It solves the #1 pain point for players — finding available grounds near them and booking without phone calls or WhatsApp negotiations. It generates direct revenue via booking commissions and drives repeat usage.",
        "highlights": [
            "Real-time slot availability — no double bookings",
            "Wallet + Razorpay dual payment options",
            "Digital Booking Pass — no need for paper receipts",
            "Smart coupon system with automatic discounts",
            "Venue reviews & ratings for trust and transparency"
        ],
        "design_direction": {
            "style": "Clean, modern UI showcase with emphasis on speed and convenience",
            "colors": "Kridaz Green (#00C187) on dark backgrounds, crisp white cards",
            "imagery": "Split-screen showing phone UI on one side, real turf/ground photo on the other. Map pins, calendar grids.",
            "formats": "Instagram carousel (swipe to see flow), single hero post, Story/Reel walkthrough",
            "cta": "\"Book Your Ground Now\" / \"Play starts here 🏏\""
        }
    },
    {
        "name": "Host a Game",
        "tagline": "Be the captain. Create the match. Invite the world.",
        "motto": "Every great match starts with someone who steps up.",
        "audience": "Organizers, team captains, social players who want to gather people",
        "category": "🏏 Player Features — Games",
        "what_it_does": "Users can host cricket matches with full customization — choose the venue (or any custom location), set date/time, define team sizes, entry fees, match format (overs), and rules. The host can invite specific players, share a public join link, or let nearby players discover and join. Once teams are filled, the host can start the match and use the Live Scoring system.",
        "user_flow": [
            "Taps 'Host Game' from the home page or bottom nav.",
            "Fills in match details: venue, date, time, team size, overs, entry fee.",
            "Sets match privacy: Public (discoverable by nearby players) or Private (invite-only).",
            "Shares the game link via WhatsApp/Instagram or waits for players to discover it.",
            "As players join, the host approves or manages the roster.",
            "On match day, taps 'Start Match' → transitions into the Live Scoring app.",
            "After the match, views analytics and shares the scorecard."
        ],
        "purpose": "This is Kridaz's social multiplier. Every hosted game brings 10–22 new users to the platform (both teams). It creates a viral loop — the host invites friends, they sign up to join, and they start hosting their own games. Revenue comes from venue bookings tied to hosted games and entry fee wallet transactions.",
        "highlights": [
            "One-tap game creation with smart defaults",
            "Public games discoverable by nearby players",
            "Built-in player invitation via link sharing",
            "Seamless transition from hosting → live scoring",
            "Entry fee management via Kridaz Wallet"
        ],
        "design_direction": {
            "style": "Energetic, action-oriented. Captain vibe — leadership, ownership.",
            "colors": "Bold green + dark theme. Accent with golden/amber for 'captain' badge feel.",
            "imagery": "Captain armband, cricket field aerial view, phone showing the game creation form, players gathering.",
            "formats": "Hero poster (\"Host Your Next Match\"), Instagram Story with countdown sticker, Reel showing the creation flow",
            "cta": "\"Host a Game Today\" / \"Be the Captain 🏆\""
        }
    },
    {
        "name": "Join Games",
        "tagline": "Short one player? Not anymore.",
        "motto": "Never miss a game because you couldn't find enough players.",
        "audience": "Individual players looking for pickup games, newcomers to an area",
        "category": "🏏 Player Features — Games",
        "what_it_does": "Players can browse publicly hosted games near their location, see match details (sport, venue, time, team size, spots remaining, entry fee), and request to join. They can filter by distance, sport, date, and available slots. Once the host approves, they get a confirmation and the match appears in their 'My Joined Games' section.",
        "user_flow": [
            "Taps 'Join Games' from the home feed or navigation.",
            "Browses a list/map of nearby public games with real-time availability.",
            "Filters by sport, date, distance, or entry fee.",
            "Taps a game card → views full details including host info, venue, rules, and who else has joined.",
            "Taps 'Join Game' → pays entry fee (if any) via Wallet.",
            "Gets a notification when the host approves.",
            "On match day, shows up and plays. Match stats are tracked via Live Scoring."
        ],
        "purpose": "This solves the biggest problem in grassroots sports: finding people to play with. For solo players or those new to a city, this is a game-changer. It also fills empty spots in hosted games, increasing completed matches and platform engagement.",
        "highlights": [
            "Discover games happening near you right now",
            "See who's already joined — play with friends or make new ones",
            "Entry fee handled via secure wallet transactions",
            "Real-time updates on spots remaining",
            "Your stats get tracked even in joined games"
        ],
        "design_direction": {
            "style": "Community-first, social, inclusive. 'Find your squad' energy.",
            "colors": "Green on dark, warm accents. Community/social vibe.",
            "imagery": "Group of players high-fiving, game cards on a phone screen, map with game pins.",
            "formats": "Instagram carousel showing the discovery → join → play flow. Social proof post (\"243 games played this week\").",
            "cta": "\"Find a Game Near You\" / \"Join the Action ⚡\""
        }
    },
    {
        "name": "Live Scoring & Match Analytics",
        "tagline": "Every ball. Every run. Every wicket. Captured forever.",
        "motto": "Turn your gully cricket match into a professional-grade scorecard.",
        "audience": "Scorers, match organizers, cricket enthusiasts, team managers",
        "category": "🏏 Player Features — Scoring",
        "what_it_does": "A full-featured ball-by-ball cricket scoring app built into Kridaz. The scorer tracks every delivery — runs, extras, wickets, boundaries — in real time. The data feeds into a rich Match Analytics dashboard showing wagon wheels, run rates, partnership charts, Manhattan charts, and individual player stats. An AI-powered commentary engine generates ball-by-ball text commentary. The match can also be live-streamed with a professional broadcast overlay.",
        "user_flow": [
            "Host starts the match → the Scoring App opens in full-screen mode.",
            "Scorer taps ball-by-ball outcomes: runs (0–6), extras (wide, no-ball), wickets (bowled, caught, etc).",
            "Live scorecard updates in real-time for all viewers.",
            "At any point, swipe to see live analytics: run rate graph, partnerships, player stats.",
            "After the match, the full analytics dashboard is available with shareable scorecard images.",
            "AI commentary generates a text recap of every ball."
        ],
        "purpose": "This is Kridaz's most differentiating feature. No other platform turns a local cricket game into a data-rich, professionally scored event. It drives deep engagement (users spend 2–3 hours per match), generates shareable content (scorecards), and creates a permanent cricket career record for every player.",
        "highlights": [
            "Ball-by-ball live scoring — works offline too",
            "Professional analytics: wagon wheels, Manhattan charts, run rate graphs",
            "AI-generated ball-by-ball commentary",
            "Shareable scorecard images for social media",
            "Player career stats tracked across all matches",
            "Live broadcast overlay for streaming (OBS compatible)"
        ],
        "design_direction": {
            "style": "Data-rich, professional, IPL broadcast quality feel. Dark theme with neon green data overlays.",
            "colors": "Dark (#050505) with neon green (#00C187) charts and white data. Stadium floodlight amber accents.",
            "imagery": "Scorecard UI screenshots, analytics dashboard, side-by-side with IPL-style scorecards. Phone showing live scoring in action.",
            "formats": "Comparison post (\"Your local match vs IPL — same analytics\"), feature deep-dive carousel, video Reel showing live scoring in action.",
            "cta": "\"Score Like a Pro\" / \"Your Match, IPL-Level Analytics 📊\""
        }
    },
    {
        "name": "Teams",
        "tagline": "Build your squad. Own your legacy.",
        "motto": "Every team starts with a dream and a roster.",
        "audience": "Team captains, friend groups, local club organizers",
        "category": "🏏 Player Features — Social",
        "what_it_does": "Users can create teams with a name, logo, and roster. They can invite existing Kridaz users or add custom (non-app) players. Teams get a digital Team Pass — a shareable card showcasing the team identity. Teams can challenge other teams, and all match stats are aggregated at the team level. Each team has a public profile page showing their record, players, and match history.",
        "user_flow": [
            "Taps 'My Teams' → 'Create Team'.",
            "Enters team name, uploads logo/banner, selects sport.",
            "Invites team members by username or share link.",
            "Team members accept → appear in the roster.",
            "Team Profile shows win/loss record, top performers, match history.",
            "Downloads or shares the Team Pass — a beautifully designed identity card."
        ],
        "purpose": "Teams create social stickiness. When someone is part of a team, their retention rate is significantly higher because they have a reason to come back (upcoming team matches, stats, rivalries). The Team Pass is a natural share vector for social media.",
        "highlights": [
            "Digital Team Pass — shareable identity card",
            "Invite by link — no need for contacts access",
            "Team-level statistics and match history",
            "Challenge other teams to matches",
            "Add custom (non-app) players to the roster"
        ],
        "design_direction": {
            "style": "Pride, identity, brotherhood. Sports jersey/crest aesthetic.",
            "colors": "Team colors (customizable) on dark. Gold accents for trophies/badges.",
            "imagery": "Team jerseys, team lineup graphic, Team Pass mockup, trophy/shield icons.",
            "formats": "\"Create Your Team\" hero post, Team Pass showcase carousel, Story template (\"Tag your squad\").",
            "cta": "\"Build Your Squad\" / \"Your Team, Your Legacy 🛡️\""
        }
    },
    {
        "name": "Find Nearby Players",
        "tagline": "Cricket is better with neighbors.",
        "motto": "Your next teammate might live two streets away.",
        "audience": "Solo players, newcomers to a city, anyone looking for playing partners",
        "category": "🏏 Player Features — Social",
        "what_it_does": "A discovery feature that shows players near your location who share your sport interests. Users can see profiles, playing preferences, skill level, and availability. They can send a connection request, start a chat, or invite them to a game directly.",
        "user_flow": [
            "Taps 'Players' from navigation.",
            "Sees a list/grid of nearby players sorted by distance.",
            "Filters by sport, skill level, or availability.",
            "Taps a player profile → sees their stats, bio, teams, and match history.",
            "Sends a connection request or invites them to an upcoming game.",
            "Starts a chat to coordinate."
        ],
        "purpose": "Solves the 'I want to play but I don't know anyone' problem. Especially powerful for people who relocate to a new city. Drives user-to-user connections which increase platform stickiness.",
        "highlights": [
            "Location-based player discovery",
            "See player stats, preferred sports, and availability",
            "One-tap game invite",
            "Direct messaging integration",
            "Privacy controls — users can hide location"
        ],
        "design_direction": {
            "style": "Social, warm, community. 'Meet your neighbor' vibe.",
            "colors": "Green + warm orange/amber accents on dark.",
            "imagery": "Map with player pins, two players shaking hands, profile cards floating on a map.",
            "formats": "Hero image (\"Find Players Near You\"), Story poll (\"Looking for a playing partner?\"), location-based promo.",
            "cta": "\"Find Players Near You\" / \"Your Next Teammate Is Nearby 📍\""
        }
    },
    {
        "name": "Find & Hire Professionals",
        "tagline": "Level up with the pros.",
        "motto": "World-class coaching and officiating, on demand.",
        "audience": "Players seeking coaching, match organizers needing umpires/scorers",
        "category": "🏏 Player Features — Professional Services",
        "what_it_does": "Users can discover and hire sports professionals — Coaches, Umpires, Scorers, Streamers, and Commentators. Browse by role, location, price, ratings, and availability. An on-demand matching system (like Uber) finds the nearest available professional in real-time using a multi-cycle dispatch algorithm with 30-second acceptance windows.",
        "user_flow": [
            "Taps 'Professionals' from navigation.",
            "Browses professionals by role, distance, and price.",
            "Taps a professional profile → sees bio, reviews, ratings, availability calendar.",
            "Books directly for a scheduled session OR uses 'Find Now' for on-demand matching.",
            "On-demand: system searches nearby online professionals, sends sequential offers.",
            "Professional accepts → user receives a confirmation with the OTP for check-in.",
            "At the venue, professional enters the OTP to start the session. Payment is handled via wallet."
        ],
        "purpose": "Creates a two-sided marketplace: professionals earn money, users get quality services. The on-demand matching (dispatch engine) is a major tech differentiator — no other cricket platform does real-time professional matching with automated payment and trust scoring.",
        "highlights": [
            "On-demand matching — find a pro in under 2 minutes",
            "Verified professionals with reviews and ratings",
            "OTP-based session verification for security",
            "Automated wallet-based payments",
            "Trust Score system ensures professional quality"
        ],
        "design_direction": {
            "style": "Premium, professional, trust-oriented. Think Uber meets sports.",
            "colors": "Deep blue or purple accents with green. Premium dark backgrounds.",
            "imagery": "Coach with whistle, umpire on field, professional profile cards, real-time matching animation.",
            "formats": "\"Hire a Coach\" feature spotlight, before/after (amateur vs coached), professional showcase carousel.",
            "cta": "\"Hire a Pro\" / \"Level Up Your Game 🎓\""
        }
    },
    {
        "name": "Kridaz Wallet",
        "tagline": "Your sports wallet. Load once, play everywhere.",
        "motto": "Fast payments, instant refunds, zero friction.",
        "audience": "All users (players, venue owners, professionals)",
        "category": "🏏 Player Features — Payments",
        "what_it_does": "A unified in-app wallet system for all transactions on Kridaz. Users can top-up via Razorpay (UPI, cards, net banking), use wallet balance for bookings, game entry fees, and professional hiring. Venue owners and professionals receive payouts to their wallet. Full transaction history with filters, withdrawal to bank account support.",
        "user_flow": [
            "Navigates to 'Wallet' from profile or bottom nav.",
            "Sees current balance, reserved balance (blocked for active bookings), and transaction history.",
            "Taps 'Add Money' → enters amount → pays via Razorpay (UPI / Card / Net Banking).",
            "Uses wallet balance for any platform transaction (bookings, games, pro hire).",
            "Views detailed transaction history with filters (date, type, status).",
            "Professionals/owners can request withdrawal to their linked bank account."
        ],
        "purpose": "The wallet is the financial backbone of the platform. It reduces payment friction (no entering card details every time), enables instant refunds for cancellations, and keeps money within the Kridaz ecosystem — driving higher transaction volumes and repeat usage.",
        "highlights": [
            "Instant top-up via UPI, cards, net banking",
            "Automatic refunds for cancellations go to wallet",
            "Reserved balance system prevents overspending",
            "Complete transaction history with export",
            "Bank withdrawal support for professionals and venue owners"
        ],
        "design_direction": {
            "style": "Fintech-clean, trust-inspiring. Banking app aesthetic with sports flair.",
            "colors": "Green for positive balances, clean whites, subtle grays. Money/coin imagery.",
            "imagery": "Wallet UI on phone, coin stack animations, UPI/card logos, transaction list view.",
            "formats": "Feature explainer (\"How Kridaz Wallet Works\"), trust post (\"Secure Payments\"), promo (\"Add ₹500, Get ₹50 bonus\").",
            "cta": "\"Load Your Wallet\" / \"Pay Instantly, Play Instantly 💰\""
        }
    },
    {
        "name": "Reels (Short Videos)",
        "tagline": "Show off your best shots. Go viral.",
        "motto": "Every player has a highlight reel.",
        "audience": "Content-creating players, cricket fans, aspiring athletes",
        "category": "🏏 Player Features — Content",
        "what_it_does": "A TikTok/Instagram Reels-style short video feed within Kridaz. Users can upload cricket-related short videos (trick shots, match highlights, training clips), browse a vertical-scroll feed, like, comment, report, and share. Creators get analytics on views, likes, and engagement. Videos are stored on Cloudflare R2 for fast global delivery.",
        "user_flow": [
            "Scrolls through the Reels feed on the home page (vertical swipe).",
            "Likes, comments, or shares reels with friends.",
            "Taps 'Upload Reel' → selects video from gallery or records.",
            "Adds caption, tags, and publishes.",
            "Views analytics: views count, likes, comments, engagement rate.",
            "Top reels are featured and boosted by the algorithm."
        ],
        "purpose": "Reels drive daily active usage and content-led growth. When users share reels outside Kridaz (WhatsApp, Instagram), it brings new users organically. It also keeps users engaged between games, reducing churn during off-days.",
        "highlights": [
            "Vertical-scroll immersive video feed",
            "Upload, like, comment, share, report",
            "Creator analytics dashboard",
            "Content moderation by admin team",
            "Fast global delivery via Cloudflare R2"
        ],
        "design_direction": {
            "style": "Bold, energetic, Gen-Z. TikTok/Reels native aesthetic.",
            "colors": "Vibrant gradients, neon green overlays on action shots.",
            "imagery": "Phone showing reel feed, cricket highlight clips, play button overlays, creator analytics.",
            "formats": "\"Upload Your First Reel\" tutorial carousel, creator spotlight, challenge announcement post.",
            "cta": "\"Upload Your Reel\" / \"Show the World Your Game 🎬\""
        }
    },
    {
        "name": "Community Feed (Posts & Stories)",
        "tagline": "Your cricket social network.",
        "motto": "Share your wins, your losses, and everything in between.",
        "audience": "All users — social engagement layer",
        "category": "🏏 Player Features — Social",
        "what_it_does": "A social media feed where users can create text/image posts, share stories (24-hour disappearing content), follow other users, like, comment, and save content. The feed shows posts from followed users and nearby community members. Users can also share match scorecards, team achievements, and venue reviews as posts.",
        "user_flow": [
            "Opens the app → lands on the Home feed.",
            "Scrolls through posts from followed users and the community.",
            "Taps '+' to create a new post (text, image, or match scorecard share).",
            "Creates a Story (photo/video that disappears in 24 hours).",
            "Likes, comments, saves, or reports posts.",
            "Follows other users to curate their feed."
        ],
        "purpose": "The social layer transforms Kridaz from a utility app into a daily-use platform. Without social, users only open the app when they want to book. With social, they open it to scroll, engage, and stay connected — dramatically increasing DAU and session time.",
        "highlights": [
            "Posts with images, text, and match scorecard embeds",
            "24-hour Stories with rich media",
            "Follow system for curated feeds",
            "Like, comment, save, share, report",
            "Nearby community discovery"
        ],
        "design_direction": {
            "style": "Social media native. Instagram-meets-sports feed aesthetic.",
            "colors": "Dark theme with green accents. Content-first with minimal chrome.",
            "imagery": "Feed scrolling UI, post creation screens, story bubbles at top.",
            "formats": "\"Share Your Cricket Life\" brand post, community highlights carousel, UGC (user-generated content) reshare templates.",
            "cta": "\"Share Your Story\" / \"Join the Cricket Community 🤝\""
        }
    },
    {
        "name": "Real-time Chat",
        "tagline": "Coordinate your matches without leaving the app.",
        "motto": "Game talk, right where it belongs.",
        "audience": "All users — communication layer",
        "category": "🏏 Player Features — Communication",
        "what_it_does": "One-on-one real-time messaging between users. Powered by WebSockets for instant message delivery. Users can chat with teammates, opponents, professionals, or venue owners. Supports text messages with read receipts and online status indicators.",
        "user_flow": [
            "Taps 'Messages' from bottom navigation.",
            "Sees list of active conversations sorted by most recent.",
            "Taps a conversation → opens chat window.",
            "Sends text messages in real-time.",
            "Starts new conversation from a player/professional profile."
        ],
        "purpose": "Chat removes the need to switch to WhatsApp for match coordination. Keeping conversations within Kridaz reduces user drop-off and ensures all match-related communication is in one place.",
        "highlights": [
            "Real-time messaging via WebSockets",
            "Online status indicators",
            "Start chat from any profile",
            "Conversation history preserved"
        ],
        "design_direction": {
            "style": "Clean, minimal messaging UI. Dark bubbles on dark background.",
            "colors": "Green message bubbles for sent, dark gray for received.",
            "imagery": "Chat UI mockup, message notification pop-up, two players chatting illustration.",
            "formats": "Feature announcement post, comparison (\"Stop switching to WhatsApp\").",
            "cta": "\"Chat with your team\" / \"Game plans start here 💬\""
        }
    },
    {
        "name": "Player Leaderboard & Career Stats",
        "tagline": "Every run counts. Every wicket matters.",
        "motto": "Build your cricket career, one match at a time.",
        "audience": "Competitive players, stats enthusiasts, aspiring cricketers",
        "category": "🏏 Player Features — Gamification",
        "what_it_does": "A global and local leaderboard ranking players by batting average, bowling economy, total runs, wickets, and overall performance. Every ball tracked in the Live Scoring system feeds into a player's career stats. Players can view their career profile showing lifetime stats, badges earned, match history, and performance trends.",
        "user_flow": [
            "Taps 'Leaderboard' from navigation.",
            "Views top players by different categories (runs, wickets, average, etc).",
            "Filters by location, time period, or sport.",
            "Taps own profile → sees complete career statistics.",
            "Shares career card on social media."
        ],
        "purpose": "Gamification is the retention engine. When players see their stats improving, they want to play more. The leaderboard creates healthy competition and gives players a reason to track every match on Kridaz rather than playing off-platform.",
        "highlights": [
            "Auto-generated career stats from match scoring",
            "Global and local leaderboards",
            "Badges and achievements system",
            "Shareable career cards for social media",
            "Performance trend tracking over time"
        ],
        "design_direction": {
            "style": "Competitive, data-driven, aspirational. Esports/fantasy sports leaderboard vibe.",
            "colors": "Gold (#FFD700) for #1, silver, bronze. Green accent. Dark premium background.",
            "imagery": "Leaderboard table UI, trophy icons, player stat cards, podium graphic.",
            "formats": "Weekly leaderboard announcement, player spotlight (\"Player of the Week\"), milestone celebration post.",
            "cta": "\"Climb the Leaderboard\" / \"Track Your Cricket Career 📈\""
        }
    },
    {
        "name": "Blogs",
        "tagline": "Cricket stories, tips, and insights.",
        "motto": "Learn from the game. Share with the community.",
        "audience": "Cricket enthusiasts, readers, content consumers",
        "category": "🏏 Player Features — Content",
        "what_it_does": "An editorial blog section where the Kridaz team and community publish cricket-related articles — match recaps, coaching tips, equipment reviews, tournament guides, and platform updates. Users can read, like, and share blog posts.",
        "user_flow": [
            "Navigates to 'Blogs' from navigation.",
            "Browses published articles by category.",
            "Taps an article → reads full content with images.",
            "Likes or shares the article."
        ],
        "purpose": "SEO-driven content brings organic traffic from Google. Blog posts also establish Kridaz as a thought leader in the grassroots cricket space. They serve as landing pages for new users discovering the platform via search.",
        "highlights": [
            "Admin-published editorial content",
            "Categorized articles for easy discovery",
            "SEO-optimized for organic traffic",
            "Shareable to social media"
        ],
        "design_direction": {
            "style": "Editorial, clean, magazine-quality.",
            "colors": "White/light backgrounds with green highlights. Clean typography.",
            "imagery": "Blog layout mockup, article cards, reading on phone, cricket action photos.",
            "formats": "Article promotion post, quote card from article, \"New on the blog\" announcement.",
            "cta": "\"Read the Latest\" / \"Cricket Insights 📰\""
        }
    },
    {
        "name": "Global Search",
        "tagline": "Find anything in one search.",
        "motto": "Venues, players, teams, games — all in one search bar.",
        "audience": "All users",
        "category": "🏏 Player Features — Navigation",
        "what_it_does": "A unified search experience that lets users find venues, players, teams, games, professionals, and blog posts from a single search bar. Results are categorized and ranked by relevance.",
        "user_flow": [
            "Taps the search icon from any page.",
            "Types a query — gets instant categorized results.",
            "Taps a result → navigates to the relevant page."
        ],
        "purpose": "Reduces friction for power users who know exactly what they're looking for. A good search experience is the mark of a mature platform.",
        "highlights": [
            "Unified search across all content types",
            "Instant results with categorization",
            "Recent searches saved"
        ],
        "design_direction": {
            "style": "Minimal, utility-focused.",
            "colors": "Clean dark with white text. Green highlight on active results.",
            "imagery": "Search bar with auto-complete dropdown, categorized results.",
            "formats": "Feature tip post (\"Did you know you can search for anything on Kridaz?\").",
            "cta": "\"Search for anything\" / \"Find it on Kridaz 🔍\""
        }
    },

    # ═══════════════════════════════════════════════════════════
    # PERSONA 2: VENUE OWNER (PARTNER)
    # ═══════════════════════════════════════════════════════════
    {
        "name": "Venue Owner Dashboard",
        "tagline": "Your venue empire, at a glance.",
        "motto": "Run your sports business smarter, not harder.",
        "audience": "Turf/ground owners, sports facility managers",
        "category": "🏟️ Venue Owner Features",
        "what_it_does": "A comprehensive business dashboard for venue owners. Shows key metrics at a glance: today's bookings, revenue (daily/weekly/monthly), occupancy rate, upcoming bookings timeline, recent reviews, and alerts. Acts as the command center for the entire venue business.",
        "user_flow": [
            "Logs in as venue owner → lands on Dashboard.",
            "Sees today's snapshot: bookings count, revenue, occupancy %.",
            "Views upcoming bookings timeline for the day.",
            "Checks recent reviews and ratings.",
            "Navigates to specific sections (bookings, revenue, turfs) from dashboard cards."
        ],
        "purpose": "The dashboard converts venue owners from 'listing a turf' to 'running a digital sports business'. The more insights they see, the more they invest in the platform (better photos, competitive pricing, promotions). This drives listing quality and bookings volume.",
        "highlights": [
            "Real-time revenue and booking metrics",
            "Occupancy rate tracking",
            "Review monitoring",
            "Quick-action cards for common tasks"
        ],
        "design_direction": {
            "style": "Professional SaaS dashboard. Clean, data-rich, business-oriented.",
            "colors": "Dark theme with green metrics, white cards, amber alerts.",
            "imagery": "Dashboard UI mockup with charts, venue owner using tablet at their ground.",
            "formats": "\"Partner with Kridaz\" hero post, benefits infographic, partner testimonial.",
            "cta": "\"List Your Venue\" / \"Grow Your Sports Business 📊\""
        }
    },
    {
        "name": "Turf Management",
        "tagline": "Add, edit, and manage your grounds with ease.",
        "motto": "Your ground, fully digitized.",
        "audience": "Venue owners with one or multiple turfs/grounds",
        "category": "🏟️ Venue Owner Features",
        "what_it_does": "Full CRUD management for sports venues. Owners can add new turfs with photos, amenities, pricing, time slots, sports types, and location. They can edit existing turf details, manage time slot availability, set dynamic pricing, and control venue status (active/draft/paused). Multi-turf support for owners with multiple venues.",
        "user_flow": [
            "Navigates to 'Turfs' → sees all owned venues.",
            "Taps 'Add Turf' → fills multi-step form: basic info, photos, pricing, time slots, amenities.",
            "Uploads high-quality venue photos (Cloudinary hosted).",
            "Sets pricing per hour, special rates for weekends/holidays.",
            "Configures available time slots.",
            "Publishes the turf → it goes live and appears in search results.",
            "Edits existing turf details anytime."
        ],
        "purpose": "Easy turf management encourages venue owners to keep their listings accurate and attractive. Better listings lead to more bookings. The multi-step form ensures complete information, which builds trust with players.",
        "highlights": [
            "Multi-step guided turf creation",
            "Photo gallery with Cloudinary CDN delivery",
            "Dynamic pricing support",
            "Multi-turf management for venue chains",
            "Draft mode before publishing"
        ],
        "design_direction": {
            "style": "Professional, clean form UI. Before/after transformation.",
            "colors": "Green success states, clean whites, structured layout.",
            "imagery": "Turf creation form mockup, beautiful venue photos, before (empty field) vs after (listed on Kridaz).",
            "formats": "\"List Your Venue in 5 Minutes\" tutorial carousel, partner onboarding guide.",
            "cta": "\"Add Your Turf\" / \"Go Digital Today 🏟️\""
        }
    },
    {
        "name": "Revenue & Banking",
        "tagline": "Track every rupee. Withdraw anytime.",
        "motto": "Transparent earnings, instant withdrawals.",
        "audience": "Venue owners focused on business financials",
        "category": "🏟️ Venue Owner Features",
        "what_it_does": "Complete financial management for venue owners. Revenue dashboard showing earnings breakdown by turf, time period, and booking type. Booking-level transaction details. Bank account linking for withdrawals. Withdrawal request system with status tracking. Revenue trend charts and export capabilities.",
        "user_flow": [
            "Navigates to 'Revenue' from partner dashboard.",
            "Views earnings summary: total revenue, this month, this week, today.",
            "Drills into revenue by turf, by day, or by booking.",
            "Navigates to 'Banking' → links bank account details.",
            "Requests withdrawal → tracks withdrawal status.",
            "Downloads revenue reports."
        ],
        "purpose": "Financial transparency is the #1 factor for venue owner retention. When owners see clear earnings and can withdraw easily, they trust the platform and stay active. This section directly impacts partner satisfaction and listing supply.",
        "highlights": [
            "Real-time revenue tracking",
            "Per-turf earnings breakdown",
            "Bank account linking for withdrawals",
            "Withdrawal request with status tracking",
            "Revenue trend charts"
        ],
        "design_direction": {
            "style": "Fintech/banking app aesthetic. Trust, transparency, professionalism.",
            "colors": "Green for earnings, blue for charts, clean white cards on dark.",
            "imagery": "Revenue dashboard UI, earnings graph going up, bank transfer illustration.",
            "formats": "\"Earn with Kridaz\" income showcase, partner earnings milestone post.",
            "cta": "\"Track Your Earnings\" / \"Your Revenue, Your Control 💰\""
        }
    },
    {
        "name": "Venue Intelligence & Promotions",
        "tagline": "Data-driven decisions for your sports business.",
        "motto": "Know your venue better than your competition.",
        "audience": "Business-savvy venue owners, venue chain operators",
        "category": "🏟️ Venue Owner Features",
        "what_it_does": "Analytics and insights dashboard showing booking patterns, peak hours, customer demographics, and review trends. Promotions module lets venue owners create and manage discount coupons, time-based offers, and marketing campaigns to drive off-peak bookings.",
        "user_flow": [
            "Navigates to 'Intelligence' → sees booking pattern heatmap.",
            "Views peak hours, most popular time slots, customer retention metrics.",
            "Identifies off-peak hours with low occupancy.",
            "Navigates to 'Promotions' → creates a coupon (% off or flat discount).",
            "Sets coupon validity period, usage limits, and applicable turfs.",
            "Shares promotion on social media to drive bookings."
        ],
        "purpose": "Intelligence features help venue owners optimize pricing and fill empty slots. Promotions during off-peak hours increase overall occupancy and revenue. This positions Kridaz as a complete business management tool, not just a listing platform.",
        "highlights": [
            "Booking pattern heatmaps and peak hour analysis",
            "Customer demographic insights",
            "Create and manage promotional coupons",
            "Off-peak optimization suggestions",
            "Review trend analysis"
        ],
        "design_direction": {
            "style": "Analytics dashboard, smart business tools. Data visualization focus.",
            "colors": "Heat map colors (cool blues to warm reds), green accents, dark background.",
            "imagery": "Heatmap charts, coupon creation UI, analytics graphs, smart suggestions.",
            "formats": "\"Smart Tools for Smart Owners\" feature carousel, data insight infographic.",
            "cta": "\"Optimize Your Venue\" / \"Fill Every Slot 📈\""
        }
    },

    # ═══════════════════════════════════════════════════════════
    # PERSONA 3: PROFESSIONAL
    # ═══════════════════════════════════════════════════════════
    {
        "name": "Professional Dashboard & On-Demand Matching",
        "tagline": "Go online. Get matched. Get paid.",
        "motto": "Your skills are in demand. Get hired instantly.",
        "audience": "Coaches, umpires, scorers, streamers, commentators, cheerleaders",
        "category": "🎓 Professional Features",
        "what_it_does": "A complete professional portal. Professionals register their role, set their hourly rate, and go 'online' to receive booking offers. The on-demand matching system works like ride-hailing: when a user requests a professional, the dispatch engine sends sequential offers to the nearest available pros with 30-second acceptance windows. Includes a dashboard with earnings, active bookings, and performance metrics.",
        "user_flow": [
            "Registers as a professional (coach/umpire/scorer/streamer).",
            "Sets hourly rate, availability, and service area.",
            "Taps 'Go Online' → starts receiving match offers.",
            "Gets a notification with match details: location, duration, payout.",
            "Accepts within 30 seconds → gets confirmed booking.",
            "Navigates to venue → enters OTP from the user to check in.",
            "Completes session → payment is auto-credited to wallet.",
            "Views earnings, ratings, and trust score on dashboard."
        ],
        "purpose": "This creates an Uber-like marketplace for sports professionals. It's a new income stream for coaches and umpires who previously relied on word-of-mouth. For Kridaz, it generates commission revenue and adds professional services as a differentiator no competitor offers.",
        "highlights": [
            "Real-time on-demand matching (like Uber for sports pros)",
            "6 professional roles supported",
            "30-second offer acceptance windows",
            "OTP-verified session check-in",
            "Automated wallet-based payments",
            "Trust Score system rewards reliability"
        ],
        "design_direction": {
            "style": "Professional, sleek, ride-hailing app aesthetic. Status indicators, earning counters.",
            "colors": "Green for 'online' status, amber for 'offer received', dark premium background.",
            "imagery": "Coach/umpire portrait with 'ONLINE' badge, offer notification card, earnings counter.",
            "formats": "\"Earn as a Cricket Pro\" recruitment post, role showcase carousel (coach, umpire, scorer, streamer), earnings testimonial.",
            "cta": "\"Register as a Pro\" / \"Start Earning Today 🎓\""
        }
    },
    {
        "name": "Trust Score System",
        "tagline": "Reliability earns rewards.",
        "motto": "Your reputation is your most valuable asset.",
        "audience": "All professionals on the platform",
        "category": "🎓 Professional Features",
        "what_it_does": "A reputation system that tracks professional reliability through a numerical Trust Score (0–100). Points are earned for accepting bookings (+1), completing sessions, and receiving positive reviews. Points are deducted for ignoring offers (-0.5 per cycle), no-shows, and negative reviews. Two consecutive skipped offers force the professional offline. Trust Score history is fully transparent.",
        "user_flow": [
            "Professional views their Trust Score on the dashboard header.",
            "Taps 'Trust Score' → sees full history of score changes.",
            "Each event shows: date, reason, points gained/lost.",
            "High trust score = priority in matching queue (shown first to users).",
            "Low trust score = lower priority and potential account review."
        ],
        "purpose": "Trust Score ensures platform quality. Without it, professionals could accept and ghost, degrading the user experience. It creates a virtuous cycle: reliable pros get more bookings → earn more → stay on the platform.",
        "highlights": [
            "Transparent scoring with full event history",
            "Rewards reliability, penalizes ghosting",
            "High score = priority matching placement",
            "Auto-offline after 2 consecutive skips",
            "Score influences earnings potential"
        ],
        "design_direction": {
            "style": "Gamification, achievement badges, progress bars.",
            "colors": "Green (high score), amber (medium), red (low). Dark background.",
            "imagery": "Trust score gauge/meter, history timeline, badge icons, shield graphic.",
            "formats": "\"How Trust Score Works\" explainer carousel, milestone badge celebration post.",
            "cta": "\"Build Your Reputation\" / \"Reliability Pays Off 🛡️\""
        }
    },
    {
        "name": "Live Streaming & Broadcast Overlay",
        "tagline": "Stream your match to the world.",
        "motto": "Turn local matches into live broadcasts.",
        "audience": "Streamers, content creators, tournament organizers",
        "category": "🎓 Professional Features",
        "what_it_does": "Integration with YouTube and Facebook for live streaming cricket matches. Streamers can set up their broadcast with RTMP keys, configure OBS settings, and use a professional scoreboard overlay that updates in real-time from the Live Scoring system. Multiple overlay themes available. The overlay is served as a browser source URL that OBS captures.",
        "user_flow": [
            "Streamer navigates to 'Stream Setup' for a specific match.",
            "Connects YouTube/Facebook account via OAuth.",
            "Configures stream: title, description, privacy settings.",
            "Copies the overlay URL → adds as a browser source in OBS.",
            "Starts streaming → overlay shows live score, batsman/bowler stats, run rate.",
            "Viewers watch on YouTube/Facebook with professional-grade graphics.",
            "After match, the VOD (video on demand) is automatically saved."
        ],
        "purpose": "Live streaming transforms local matches into content events. It attracts an audience beyond the players, drives brand visibility, and positions Kridaz as a professional cricket production platform. Tournament organizers see massive value in this for their events.",
        "highlights": [
            "YouTube + Facebook live integration",
            "Professional OBS overlay with real-time scoring",
            "Multiple overlay themes (IPL-style, minimal, classic)",
            "One-click stream setup",
            "Automatic VOD archiving"
        ],
        "design_direction": {
            "style": "Broadcast/TV production aesthetic. IPL-quality overlay showcase.",
            "colors": "Rich dark with neon green data, broadcast frame styling.",
            "imagery": "OBS overlay screenshot, YouTube live embed, streamer setup, split-screen (real match + overlay).",
            "formats": "Overlay theme showcase carousel, \"Stream Your Match\" tutorial, live broadcast announcement template.",
            "cta": "\"Go Live\" / \"Broadcast Your Match 📺\""
        }
    },
]


# ──────────────────────────────────────────────────────────────
# DOCUMENT GENERATION
# ──────────────────────────────────────────────────────────────
def generate_document():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # ── COVER PAGE ──
    for _ in range(6):
        doc.add_paragraph("")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("KRIDAZ")
    run.font.size = Pt(48)
    run.font.color.rgb = KRIDAZ_GREEN
    run.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Platform Feature Documentation")
    run.font.size = Pt(24)
    run.font.color.rgb = DARK_TEXT

    doc.add_paragraph("")

    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = desc.add_run("For Marketing & Design Team")
    run.font.size = Pt(14)
    run.font.color.rgb = MEDIUM_GRAY

    doc.add_paragraph("")

    version = doc.add_paragraph()
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = version.add_run("Version 1.0 — June 2026")
    run.font.size = Pt(11)
    run.font.color.rgb = MEDIUM_GRAY

    doc.add_paragraph("")

    tagline = doc.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tagline.add_run("\"The Ultimate Cricket Platform\"")
    run.font.size = Pt(14)
    run.font.color.rgb = KRIDAZ_GREEN
    run.italic = True

    doc.add_page_break()

    # ── TABLE OF CONTENTS (Manual) ──
    toc_heading = doc.add_heading("Table of Contents", level=1)
    for run in toc_heading.runs:
        run.font.color.rgb = DARK_TEXT

    doc.add_paragraph("")

    # Group features by category prefix
    categories = {}
    for i, f in enumerate(FEATURES, 1):
        cat = f["category"]
        if cat not in categories:
            categories[cat] = []
        categories[cat].append((i, f["name"]))

    for cat, features in categories.items():
        p = doc.add_paragraph()
        run = p.add_run(cat)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = DARK_TEXT

        for num, name in features:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(f"#{num}  {name}")
            run.font.size = Pt(10)
            run.font.color.rgb = MEDIUM_GRAY

    doc.add_page_break()

    # ── INTRODUCTION ──
    intro_heading = doc.add_heading("About This Document", level=1)
    for run in intro_heading.runs:
        run.font.color.rgb = DARK_TEXT

    intro_text = (
        "This document provides a comprehensive overview of every feature available on the Kridaz platform. "
        "It is designed for the marketing and design team to understand what each feature does, who it serves, "
        "how users interact with it, and what visual direction to take when creating marketing graphics and social media posts.\n\n"
        "Each feature includes:\n"
        "• A marketing tagline and emotional motto\n"
        "• A plain-English description of what it does\n"
        "• A step-by-step user flow showing how the feature is used\n"
        "• The business purpose and value proposition\n"
        "• Key highlights for marketing copy\n"
        "• Visual design direction for graphic posts (style, colors, imagery, formats, CTA)\n\n"
        "The features are organized by user persona:\n"
        "🏏 Player (User) — 14 features\n"
        "🏟️ Venue Owner (Partner) — 4 feature groups\n"
        "🎓 Professional (Coach/Umpire/Scorer/Streamer) — 3 feature groups\n\n"
        "Total: 21 documented features covering every major capability of the platform."
    )

    p = doc.add_paragraph(intro_text)
    for run in p.runs:
        run.font.size = Pt(10)

    doc.add_page_break()

    # ── FEATURE SECTIONS ──
    current_category = None
    for i, feature in enumerate(FEATURES, 1):
        cat = feature["category"]

        # Add category header if new category
        cat_prefix = cat.split("—")[0].strip() if "—" in cat else cat
        if cat_prefix != current_category:
            current_category = cat_prefix
            cat_heading = doc.add_heading(cat_prefix, level=1)
            for run in cat_heading.runs:
                run.font.color.rgb = DARK_TEXT
                run.font.size = Pt(22)
            add_horizontal_rule(doc)

        add_feature_section(doc, i, feature)

    # ── CLOSING PAGE ──
    doc.add_page_break()
    for _ in range(8):
        doc.add_paragraph("")

    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = closing.add_run("— End of Document —")
    run.font.size = Pt(14)
    run.font.color.rgb = MEDIUM_GRAY

    doc.add_paragraph("")

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = contact.add_run("For questions about any feature, contact the product team.\nKridaz © 2026")
    run.font.size = Pt(10)
    run.font.color.rgb = MEDIUM_GRAY

    # ── SAVE ──
    output_path = os.path.join(os.path.dirname(__file__), "Kridaz_Marketing_Feature_Documentation.docx")
    doc.save(output_path)
    print(f"Document saved to: {output_path}")
    return output_path


if __name__ == "__main__":
    generate_document()
